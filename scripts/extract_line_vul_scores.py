#!/usr/bin/env python3
"""Extract workbook VIC/VFC occurrences and their Hugging Face test scores."""
import argparse
import csv
from collections import Counter
from pathlib import Path
import re
import sys
try:
    from .check_prediction_labels import MATCH_COLUMN, prediction_matches_label
except ImportError:
    from check_prediction_labels import MATCH_COLUMN, prediction_matches_label

SHA = re.compile(r"[0-9a-f]{7,40}")
FIELDS = ['dataset', 'excel_row', 'commit_prefix', 'vul_lines', 'model',
          'config', 'seed', 'run', 'budget', 'label', 'prediction', 'probability', MATCH_COLUMN]


def read_commits(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(v).strip() if v is not None else '' for v in next(rows)]
            if not {'Vul_commit', 'Patch_commit'} <= set(headers):
                raise ValueError(f'{sheet.title}: missing Vul_commit/Patch_commit columns')
            for row_number, values in enumerate(rows, 2):
                row = dict(zip(headers, values))
                for column, kind in [('Vul_commit', 'VIC'), ('Patch_commit', 'VFC')]:
                    value = row.get(column)
                    if value is None or not str(value).strip():
                        continue
                    prefix = str(value).strip().lower()
                    if not SHA.fullmatch(prefix):
                        raise ValueError(f'{sheet.title}!{column}, row {row_number}: invalid SHA {value!r}')
                    records.append(dict(dataset=sheet.title.lower(), sheet=sheet.title,
                                        excel_row=row_number, commit_type=kind,
                                        commit_prefix=prefix, cve=row.get('CVE') or '',
                                        vul_lines=row.get('Vul_lines') or ''))
    finally:
        book.close()
    if not records:
        raise ValueError('No commits found in workbook')
    return records


def parse_path(path):
    """Read legacy output paths and paths with a sampling-mode directory."""
    parts = path.split('/')
    if len(parts) == 7 and parts[3] in {'sampling', 'no_sampling'}:
        parts = parts[:3] + parts[4:]
    if len(parts) != 6 or parts[0] != 'output':
        return None
    _, dataset, model, config, experiment, filename = parts
    run = re.fullmatch(r'(?:seed_(\d+|default)_)?run_(\d+)', experiment)
    budget = re.fullmatch(re.escape(model) + r'_budget_(.+)_test_scores\.csv', filename)
    if not run or not budget:
        return None
    return dict(dataset=dataset, model=model, config=config, seed=run[1] or 'default',
                run=run[2], budget=budget[1].replace('p', '.'))


def extract(records, score_path, metadata, source, revision):
    with open(score_path, newline='', encoding='utf-8-sig') as handle:
        reader = csv.DictReader(handle)
        if not {'commit_id', 'label', 'prediction', 'probability'} <= set(reader.fieldnames or []):
            raise ValueError(f'{source}: missing required score columns')
        scores = {}
        for row in reader:
            sha = row['commit_id'].strip().lower()
            if not SHA.fullmatch(sha):
                raise ValueError(f'{source}: invalid commit_id {sha!r}')
            if sha in scores:
                raise ValueError(f'{source}: duplicate commit_id {sha}')
            scores[sha] = row
    for record in records:
        if record['dataset'] != metadata['dataset']:
            continue
        matches = [sha for sha in scores if sha.startswith(record['commit_prefix'])]
        result = {**record, **metadata, 'source_file': source, 'revision': revision}
        result['status'] = 'matched' if len(matches) == 1 else 'missing' if not matches else 'ambiguous'
        if len(matches) == 1:
            result.update({key: scores[matches[0]][key] for key in
                           ['commit_id', 'label', 'prediction', 'probability']})
        result[MATCH_COLUMN] = prediction_matches_label(result)
        yield result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=Path(__file__).resolve().parents[1] / 'Data_line_vul.xlsx')
    parser.add_argument('--repo-id', default='TheSyx/vulguard_lite')
    parser.add_argument('--revision', default='main')
    parser.add_argument('--output', type=Path, default=Path('line_vul_scores.csv'))
    parser.add_argument('--local-root', type=Path, help='Read an existing output/ tree here instead of Hugging Face')
    parser.add_argument('--cache-dir', type=Path)
    parser.add_argument('--files', nargs='+', help='Exact repository paths; skip repository listing')
    for name in ['datasets', 'models', 'configs', 'seeds', 'runs', 'budgets']:
        parser.add_argument('--' + name, nargs='+', help='Optional exact filters (space separated)')
    args = parser.parse_args()
    records = read_commits(args.input)
    print(f'Workbook: {len(records)} occurrences; ' + str(dict(Counter(r['commit_type'] for r in records))), file=sys.stderr)
    print(f"Unique dataset/SHA pairs: {len({(r['dataset'], r['commit_prefix']) for r in records})}", file=sys.stderr)
    revision = args.revision
    if args.local_root:
        paths = args.files or [p.relative_to(args.local_root).as_posix()
                               for p in args.local_root.glob('output/**/*_test_scores.csv')]
    else:
        from huggingface_hub import HfApi, hf_hub_download
        api = HfApi()
        # Pin all downloads to a single immutable dataset version.
        revision = api.repo_info(args.repo_id, repo_type='dataset', revision=revision).sha
        paths = args.files or api.list_repo_files(args.repo_id, repo_type='dataset', revision=revision)
    selected = []
    for path in sorted(set(paths)):
        meta = parse_path(path)
        if meta is None:
            if args.files:
                raise ValueError(f'Unsupported score path: {path}')
            continue
        if meta['dataset'] not in {r['dataset'] for r in records}:
            continue
        if any(getattr(args, plural) and meta[singular] not in getattr(args, plural)
               for plural, singular in [('datasets', 'dataset'), ('models', 'model'),
                                        ('configs', 'config'), ('seeds', 'seed'),
                                        ('runs', 'run'), ('budgets', 'budget')]):
            continue
        selected.append((path, meta))
    if not selected:
        raise ValueError('No test score files match the selected filters')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.output.with_suffix(args.output.suffix + '.tmp')
    counts = Counter()
    with temporary.open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        for index, (path, meta) in enumerate(selected, 1):
            print(f'[{index}/{len(selected)}] {path}', file=sys.stderr)
            local = args.local_root / path if args.local_root else hf_hub_download(
                repo_id=args.repo_id, repo_type='dataset', revision=revision,
                filename=path, cache_dir=args.cache_dir)
            for result in extract(records, local, meta, path, revision):
                writer.writerow({field: result.get(field, '') for field in FIELDS})
                counts[result['status']] += 1
    temporary.replace(args.output)
    covered = {meta['dataset'] for _, meta in selected}
    absent = {r['dataset'] for r in records} - covered
    if absent:
        print(f'No selected results for workbook datasets: {sorted(absent)}', file=sys.stderr)
    print(f'Written {args.output}: {dict(counts)}', file=sys.stderr)


if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        sys.exit(f'Error: {exc}')
